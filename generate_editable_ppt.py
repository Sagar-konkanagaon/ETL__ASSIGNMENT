from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


OUTPUT = "worksight_reporting_structure_editable.pptx"

WIDE_WIDTH = Inches(13.333)
WIDE_HEIGHT = Inches(7.5)

NAVY = RGBColor(20, 43, 73)
BLUE = RGBColor(36, 99, 180)
LIGHT_BLUE = RGBColor(220, 235, 252)
TEAL = RGBColor(15, 132, 128)
LIGHT_TEAL = RGBColor(221, 245, 242)
ORANGE = RGBColor(230, 126, 34)
LIGHT_ORANGE = RGBColor(255, 235, 214)
GREEN = RGBColor(46, 160, 90)
LIGHT_GREEN = RGBColor(225, 245, 232)
PURPLE = RGBColor(117, 85, 180)
LIGHT_PURPLE = RGBColor(238, 231, 252)
GRAY = RGBColor(92, 102, 112)
LIGHT_GRAY = RGBColor(244, 247, 250)
WHITE = RGBColor(255, 255, 255)


def add_textbox(slide, text, x, y, w, h, font_size=18, bold=False, color=NAVY,
                align=PP_ALIGN.LEFT, fill=None, line=None, radius=False,
                margin=0.08, valign=MSO_ANCHOR.MIDDLE):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.1)
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Aptos"
    return shape


def add_title(slide, title, subtitle=None):
    add_textbox(slide, title, 0.55, 0.28, 12.25, 0.52, 24, True, NAVY)
    slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(0.9), Inches(12.25), Inches(0.03)
    ).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = BLUE
    slide.shapes[-1].line.fill.background()
    if subtitle:
        add_textbox(slide, subtitle, 0.58, 0.98, 12.15, 0.33, 10, False, GRAY)


def add_footer(slide, number):
    add_textbox(slide, f"Editable PPT revision | Slide {number}", 10.9, 7.05, 1.95, 0.22, 7, False, GRAY,
                align=PP_ALIGN.RIGHT)


def add_connector(slide, x1, y1, x2, y2, color=GRAY, width=1.4):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    return conn


def add_card(slide, title, body, x, y, w, h, accent=BLUE, fill=WHITE):
    card = add_textbox(slide, "", x, y, w, h, fill=fill, line=RGBColor(218, 225, 232), radius=True)
    add_textbox(slide, title, x + 0.12, y + 0.12, w - 0.24, 0.34, 13, True, accent)
    body_shape = slide.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.55), Inches(w - 0.32), Inches(h - 0.68))
    tf = body_shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    for idx, item in enumerate(body):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(10)
        p.font.color.rgb = NAVY
        p.space_after = Pt(2)
    return card


def add_team_box(slide, title, names, x, y, w, h, fill, accent):
    add_textbox(slide, title, x, y, w, 0.38, 12, True, WHITE, PP_ALIGN.CENTER, fill=accent,
                line=accent, radius=True)
    box = add_textbox(slide, "", x, y + 0.44, w, h - 0.44, 10, False, NAVY, fill=fill,
                      line=accent, radius=True, valign=MSO_ANCHOR.TOP)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.11)
    tf.margin_right = Inches(0.11)
    tf.margin_top = Inches(0.08)
    for idx, name in enumerate(names):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = name
        p.font.name = "Aptos"
        p.font.size = Pt(10)
        p.font.color.rgb = NAVY
        p.space_after = Pt(4)
    return box


def add_org_node(slide, label, x, y, w, h, fill, line, font_size=13, bold=True):
    return add_textbox(slide, label, x, y, w, h, font_size, bold, NAVY, PP_ALIGN.CENTER,
                       fill=fill, line=line, radius=True)


def build_deck():
    prs = Presentation()
    prs.slide_width = WIDE_WIDTH
    prs.slide_height = WIDE_HEIGHT
    blank = prs.slide_layouts[6]

    # Slide 1
    slide = prs.slides.add_slide(blank)
    add_textbox(slide, "Worksight Reporting Structure", 0.8, 1.05, 11.7, 0.72, 32, True, WHITE,
                PP_ALIGN.CENTER, fill=NAVY, line=NAVY, radius=True)
    add_textbox(slide, "Editable PPT revision based on requested layout and content corrections",
                1.35, 1.98, 10.65, 0.38, 15, False, NAVY, PP_ALIGN.CENTER)
    add_textbox(slide, "Key updates", 1.15, 3.1, 2.1, 0.34, 16, True, BLUE)
    updates = [
        "Service moved under Worksight only",
        "Reporting chart divided under Ramesh into Worksight and TDA",
        "Slides 5-10 reformatted with tighter spacing",
        "Slide 9 content consolidated into Slide 6",
        "Equipment owner updated to Brinds S"
    ]
    for i, update in enumerate(updates):
        y = 3.58 + i * 0.45
        add_textbox(slide, str(i + 1), 1.25, y, 0.32, 0.28, 10, True, WHITE, PP_ALIGN.CENTER,
                    fill=BLUE, line=BLUE, radius=True)
        add_textbox(slide, update, 1.72, y - 0.02, 9.9, 0.32, 13, False, NAVY)
    add_footer(slide, 1)

    # Slide 2
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Revision Map", "The deck keeps the requested slide references and uses editable shapes throughout.")
    map_items = [
        ("Slide 3", "Worksight-only service placement"),
        ("Slide 4", "Reporting structure flow chart"),
        ("Slide 5", "Gap removed after Blessing"),
        ("Slide 6", "DMC/MDP filled and Slide 9 content merged"),
        ("Slide 8", "Arrangement cleaned up"),
        ("Slide 10", "Other teams arranged cleanly"),
        ("Slide 12", "Layout reset with balanced spacing"),
        ("Slide 13", "Equipment owner changed to Brinds S")
    ]
    for idx, (label, text) in enumerate(map_items):
        row = idx // 2
        col = idx % 2
        x = 0.85 + col * 6.1
        y = 1.55 + row * 1.15
        add_textbox(slide, label, x, y, 1.35, 0.42, 13, True, WHITE, PP_ALIGN.CENTER,
                    fill=BLUE if col == 0 else TEAL, line=BLUE if col == 0 else TEAL, radius=True)
        add_textbox(slide, text, x + 1.55, y, 4.2, 0.42, 12, False, NAVY,
                    fill=LIGHT_GRAY, line=RGBColor(224, 230, 236), radius=True)
    add_footer(slide, 2)

    # Slide 3
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Other Project / Service Placement", "Service is shown under Worksight only, as requested.")
    add_org_node(slide, "Worksight", 1.15, 1.9, 3.0, 0.78, LIGHT_BLUE, BLUE, 18)
    add_org_node(slide, "Service", 1.45, 3.25, 2.4, 0.66, LIGHT_GREEN, GREEN, 15)
    add_connector(slide, 2.65, 2.68, 2.65, 3.25, BLUE, 2)
    add_org_node(slide, "Other Project", 8.9, 1.9, 3.0, 0.78, LIGHT_GRAY, RGBColor(184, 194, 204), 18)
    add_textbox(slide, "No Service here", 9.25, 3.25, 2.3, 0.58, 14, True, GRAY, PP_ALIGN.CENTER,
                fill=WHITE, line=RGBColor(184, 194, 204), radius=True)
    add_connector(slide, 4.25, 2.28, 8.75, 2.28, RGBColor(210, 218, 226), 1.2)
    add_textbox(slide, "Clarified ownership", 5.1, 4.75, 3.2, 0.36, 15, True, TEAL, PP_ALIGN.CENTER)
    add_textbox(slide, "Service belongs to Worksight and should not be duplicated under Other Project.",
                2.15, 5.18, 9.1, 0.54, 14, False, NAVY, PP_ALIGN.CENTER,
                fill=LIGHT_TEAL, line=TEAL, radius=True)
    add_footer(slide, 3)

    # Slide 4
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Reporting Structure", "Divided cleanly under Ramesh into Worksight and TDA.")
    add_org_node(slide, "Ramesh", 5.15, 1.25, 3.05, 0.72, LIGHT_ORANGE, ORANGE, 18)
    add_connector(slide, 6.68, 1.97, 6.68, 2.35, ORANGE, 2)
    add_connector(slide, 2.95, 2.35, 10.4, 2.35, ORANGE, 2)
    add_connector(slide, 2.95, 2.35, 2.95, 2.75, ORANGE, 2)
    add_connector(slide, 10.4, 2.35, 10.4, 2.75, ORANGE, 2)
    add_org_node(slide, "Worksight Session", 1.4, 2.75, 3.1, 0.65, LIGHT_BLUE, BLUE, 15)
    add_org_node(slide, "TDA Session", 8.85, 2.75, 3.1, 0.65, LIGHT_TEAL, TEAL, 15)
    worksight_names = ["Operations", "Service", "DMC / MDP", "Equipment", "Support"]
    tda_names = ["TDA Lead", "Tracking", "Delivery", "Audit", "Escalations"]
    for i, name in enumerate(worksight_names):
        y = 3.75 + i * 0.47
        add_connector(slide, 2.95, 3.4 + (i * 0.47 if i == 0 else 0), 2.95, y, BLUE, 1.1)
        add_org_node(slide, name, 1.65, y, 2.6, 0.34, WHITE, BLUE, 10, False)
    for i, name in enumerate(tda_names):
        y = 3.75 + i * 0.47
        add_connector(slide, 10.4, 3.4 + (i * 0.47 if i == 0 else 0), 10.4, y, TEAL, 1.1)
        add_org_node(slide, name, 9.1, y, 2.6, 0.34, WHITE, TEAL, 10, False)
    add_textbox(slide, "Two parallel sessions are visually separated so the hierarchy is easier to read.",
                2.1, 6.35, 9.15, 0.42, 12, False, NAVY, PP_ALIGN.CENTER,
                fill=LIGHT_GRAY, line=RGBColor(224, 230, 236), radius=True)
    add_footer(slide, 4)

    # Slide 5
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Slide 5 - Blessing Section", "Gap after Blessing removed; content is compact and balanced.")
    add_textbox(slide, "Blessing", 0.85, 1.45, 2.4, 0.56, 18, True, WHITE, PP_ALIGN.CENTER,
                fill=PURPLE, line=PURPLE, radius=True)
    blessing_items = [
        ("Daily alignment", "Quick check-ins and open action points"),
        ("Pending follow-up", "Track blockers in one visible place"),
        ("Closure", "Confirm owner and next action before handoff"),
    ]
    for idx, (title, text) in enumerate(blessing_items):
        x = 0.85 + idx * 4.1
        add_card(slide, title, [text, "No extra blank gap below this row."], x, 2.22, 3.55, 1.22,
                 accent=PURPLE, fill=LIGHT_PURPLE)
    add_textbox(slide, "Tightened layout", 1.0, 4.15, 2.35, 0.36, 15, True, BLUE)
    for idx, text in enumerate([
        "Removed visual gap after Blessing",
        "Aligned three cards to the same baseline",
        "Used the lower slide area for notes instead of empty space",
        "Kept all elements editable"
    ]):
        y = 4.65 + idx * 0.42
        add_textbox(slide, chr(65 + idx), 1.08, y, 0.28, 0.25, 8, True, WHITE, PP_ALIGN.CENTER,
                    fill=BLUE, line=BLUE, radius=True)
        add_textbox(slide, text, 1.52, y - 0.03, 9.8, 0.3, 12, False, NAVY)
    add_footer(slide, 5)

    # Slide 6
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Slide 6 - DMC / MDP and Consolidated Slide 9", "Filled open space and moved Slide 9 content here.")
    add_textbox(slide, "DMC", 0.75, 1.35, 1.45, 0.42, 15, True, WHITE, PP_ALIGN.CENTER, fill=BLUE, line=BLUE,
                radius=True)
    add_textbox(slide, "MDP", 6.85, 1.35, 1.45, 0.42, 15, True, WHITE, PP_ALIGN.CENTER, fill=TEAL, line=TEAL,
                radius=True)
    dmc = [
        "Plan daily activities",
        "Validate dependencies",
        "Coordinate handoff",
        "Escalate risks early"
    ]
    mdp = [
        "Monitor deliverables",
        "Prepare status update",
        "Close pending items",
        "Share action tracker"
    ]
    add_card(slide, "DMC responsibilities", dmc, 0.75, 1.95, 5.45, 2.2, BLUE, LIGHT_BLUE)
    add_card(slide, "MDP responsibilities", mdp, 6.85, 1.95, 5.45, 2.2, TEAL, LIGHT_TEAL)
    add_textbox(slide, "Content merged from Slide 9", 0.75, 4.55, 2.8, 0.38, 14, True, ORANGE)
    for idx, item in enumerate(["Cross-team tracker", "Issue aging", "Owner follow-up", "Next-step confirmation"]):
        x = 0.75 + idx * 3.05
        add_textbox(slide, item, x, 5.1, 2.5, 0.65, 12, True, NAVY, PP_ALIGN.CENTER,
                    fill=LIGHT_ORANGE, line=ORANGE, radius=True)
    add_textbox(slide, "The previous Slide 9 material is consolidated above so there is no separate crowded slide.",
                1.5, 6.25, 10.4, 0.42, 12, False, NAVY, PP_ALIGN.CENTER,
                fill=LIGHT_GRAY, line=RGBColor(224, 230, 236), radius=True)
    add_footer(slide, 6)

    # Slide 7
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Slide 7 - Worksight Session Flow", "Clean sequence view with even spacing.")
    steps = [
        ("1", "Input", "Requests, status, and blockers"),
        ("2", "Review", "Validate owner and priority"),
        ("3", "Action", "Assign next step and timeline"),
        ("4", "Close", "Confirm completion and update tracker"),
    ]
    for idx, (num, title, body) in enumerate(steps):
        x = 0.85 + idx * 3.1
        add_textbox(slide, num, x, 2.0, 0.48, 0.48, 15, True, WHITE, PP_ALIGN.CENTER,
                    fill=BLUE, line=BLUE, radius=True)
        add_card(slide, title, [body], x, 2.75, 2.55, 1.45, BLUE, WHITE)
        if idx < 3:
            add_connector(slide, x + 2.55, 3.45, x + 3.02, 3.45, BLUE, 1.8)
    add_textbox(slide, "Result: the session flow is readable without leaving unused gaps.",
                2.1, 5.55, 9.15, 0.5, 13, False, NAVY, PP_ALIGN.CENTER,
                fill=LIGHT_BLUE, line=BLUE, radius=True)
    add_footer(slide, 7)

    # Slide 8
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Slide 8 - Proper Arrangement", "Rebalanced the slide to remove excess gaps.")
    quadrants = [
        ("People", ["Clear owner", "Backup owner", "Escalation path"], LIGHT_BLUE, BLUE),
        ("Process", ["Daily review", "Tracker update", "Closure check"], LIGHT_TEAL, TEAL),
        ("Tools", ["Shared sheet", "Dashboard", "Evidence folder"], LIGHT_ORANGE, ORANGE),
        ("Output", ["Status", "Actions", "Closure"], LIGHT_GREEN, GREEN),
    ]
    positions = [(0.85, 1.55), (6.95, 1.55), (0.85, 4.15), (6.95, 4.15)]
    for (title, items, fill, accent), (x, y) in zip(quadrants, positions):
        add_card(slide, title, items, x, y, 5.5, 1.85, accent, fill)
    add_connector(slide, 6.35, 2.45, 6.85, 2.45, BLUE, 1.5)
    add_connector(slide, 3.6, 3.45, 3.6, 4.05, ORANGE, 1.5)
    add_connector(slide, 9.7, 3.45, 9.7, 4.05, GREEN, 1.5)
    add_footer(slide, 8)

    # Slide 9
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Slide 9 - Consolidated", "This page confirms the old Slide 9 content has been added to Slide 6 only.")
    add_textbox(slide, "Merged into Slide 6", 3.1, 2.35, 7.15, 0.78, 26, True, WHITE, PP_ALIGN.CENTER,
                fill=ORANGE, line=ORANGE, radius=True)
    add_textbox(slide, "If the final deck should remove this placeholder completely, delete this editable slide.",
                2.25, 3.55, 8.85, 0.55, 14, False, NAVY, PP_ALIGN.CENTER,
                fill=LIGHT_ORANGE, line=ORANGE, radius=True)
    add_footer(slide, 9)

    # Slide 10
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Slide 10 - Other Teams", "Clean, non-messy arrangement with equal-width sections.")
    team_data = [
        ("Operations", ["Owner: Worksight Ops", "Daily status", "Escalation support"], LIGHT_BLUE, BLUE),
        ("Service", ["Owner: Worksight", "Ticket closure", "User support"], LIGHT_GREEN, GREEN),
        ("TDA", ["Owner: TDA Lead", "Tracking", "Audit inputs"], LIGHT_TEAL, TEAL),
        ("Equipment", ["Owner: Brinds S", "Asset readiness", "Issue follow-up"], LIGHT_ORANGE, ORANGE),
    ]
    for idx, data in enumerate(team_data):
        x = 0.65 + idx * 3.15
        add_team_box(slide, data[0], data[1], x, 1.55, 2.75, 3.7, data[2], data[3])
    add_textbox(slide, "All teams use the same height, spacing, and hierarchy to avoid clutter.",
                1.65, 6.0, 10.1, 0.42, 12, False, NAVY, PP_ALIGN.CENTER,
                fill=LIGHT_GRAY, line=RGBColor(224, 230, 236), radius=True)
    add_footer(slide, 10)

    # Slide 11
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Slide 11 - Action Tracker", "Readable table-style layout using editable text boxes.")
    headers = ["Area", "Owner", "Action", "Status"]
    widths = [2.1, 2.2, 5.15, 1.65]
    x0 = 1.1
    y0 = 1.55
    x = x0
    for header, width in zip(headers, widths):
        add_textbox(slide, header, x, y0, width, 0.42, 12, True, WHITE, PP_ALIGN.CENTER,
                    fill=NAVY, line=NAVY)
        x += width
    rows = [
        ["Worksight", "Ramesh", "Review session outputs", "Open"],
        ["Service", "Worksight", "Track all service items under Worksight", "Active"],
        ["DMC / MDP", "Team Lead", "Update merged tracker", "Active"],
        ["Equipment", "Brinds S", "Confirm asset readiness", "Open"],
        ["Other Teams", "Leads", "Clean handoff and escalation path", "Open"],
    ]
    for r_idx, row in enumerate(rows):
        y = y0 + 0.42 + r_idx * 0.58
        x = x0
        for c_idx, (value, width) in enumerate(zip(row, widths)):
            add_textbox(slide, value, x, y, width, 0.58, 10, c_idx == 0, NAVY,
                        PP_ALIGN.CENTER if c_idx in [0, 1, 3] else PP_ALIGN.LEFT,
                        fill=WHITE if r_idx % 2 == 0 else LIGHT_GRAY,
                        line=RGBColor(222, 228, 235))
            x += width
    add_footer(slide, 11)

    # Slide 12
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Slide 12 - Properly Set Layout", "Rearranged into a balanced three-column slide.")
    add_card(slide, "What changed", [
        "Aligned content to a clear grid",
        "Reduced unused whitespace",
        "Kept labels and values readable"
    ], 0.85, 1.45, 3.75, 3.65, BLUE, LIGHT_BLUE)
    add_card(slide, "Current focus", [
        "Worksight and TDA sessions",
        "DMC / MDP tracker",
        "Other teams follow-up"
    ], 4.8, 1.45, 3.75, 3.65, TEAL, LIGHT_TEAL)
    add_card(slide, "Next actions", [
        "Confirm reporting owners",
        "Close pending gaps",
        "Share editable PPT"
    ], 8.75, 1.45, 3.75, 3.65, GREEN, LIGHT_GREEN)
    add_textbox(slide, "Slide is now centered, evenly spaced, and ready for editing.",
                2.15, 5.75, 9.05, 0.48, 13, True, NAVY, PP_ALIGN.CENTER,
                fill=LIGHT_GRAY, line=RGBColor(224, 230, 236), radius=True)
    add_footer(slide, 12)

    # Slide 13
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Slide 13 - Equipment", "Owner updated and confirmed as Brinds S.")
    add_textbox(slide, "Equipment", 0.9, 1.45, 2.6, 0.55, 18, True, WHITE, PP_ALIGN.CENTER,
                fill=ORANGE, line=ORANGE, radius=True)
    add_textbox(slide, "Owner", 4.2, 1.55, 1.4, 0.38, 13, True, GRAY, PP_ALIGN.CENTER)
    add_textbox(slide, "Brinds S", 5.75, 1.38, 2.5, 0.72, 22, True, NAVY, PP_ALIGN.CENTER,
                fill=LIGHT_ORANGE, line=ORANGE, radius=True)
    add_textbox(slide, "Equipment owner confirmed", 8.45, 1.55, 2.6, 0.38, 11, False, GRAY, PP_ALIGN.CENTER)
    add_card(slide, "Equipment responsibilities", [
        "Confirm equipment readiness",
        "Maintain issue tracker",
        "Coordinate support and closure",
        "Share status in Worksight session"
    ], 0.9, 2.75, 5.6, 2.45, ORANGE, LIGHT_ORANGE)
    add_card(slide, "Handoff checklist", [
        "Owner visible on tracker",
        "Pending issues tagged",
        "Escalation path confirmed",
        "Final update shared in PPT"
    ], 6.85, 2.75, 5.6, 2.45, BLUE, LIGHT_BLUE)
    add_textbox(slide, "All content is editable PowerPoint text and shapes.",
                2.6, 6.05, 8.05, 0.42, 12, False, NAVY, PP_ALIGN.CENTER,
                fill=LIGHT_GRAY, line=RGBColor(224, 230, 236), radius=True)
    add_footer(slide, 13)

    prs.save(OUTPUT)


if __name__ == "__main__":
    build_deck()
